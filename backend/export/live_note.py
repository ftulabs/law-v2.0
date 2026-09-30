"""The live-test short note, written into the organisers' own Word template.

The committee hands out `Live_Test_Short_Note_TEMPLATE.docx` (a copy is kept next to this file)
and expects it back printed and signed. A document that merely follows its headings makes a
steward hunt for each answer, so this fills THEIR file, cell by cell.

Every field is derived from the two runs; nothing asks the operator to copy a number:

  * section 1 — team, task, the engines ACTUALLY used (read off each run, not off the picker),
    the machine and the time the note was made;
  * section 2 — provisions exported, how many of those are absent from the 2025 baseline,
    documents fetched, minutes and cost, per engine;
  * sections 3–5 — drafted from what the runs recorded: what the logs reported as a failure,
    whether the second pass stayed off the network, where the engines disagreed, which
    indicators came back empty and which exported rows sit lowest. Each sentence states a
    fact of THIS run with its number. When the runs recorded no failure the note says so in
    those words rather than inventing a weakness — and still names what a human did not check.

The operator can edit sections 3–5 before downloading (they are the team's account, and the
template says so); section 6 records whether anything in the evidence was typed by hand.
Pure functions over the live-test state, so the screen, the .docx and the tests share one
implementation.
"""
from __future__ import annotations

import io
import os
import platform
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from ..rdtii.codes import to_rdtii_code
from ..schemas import PLACEHOLDER_LAW_NAMES
from . import engine_compare

TEMPLATE = Path(__file__).resolve().parent / "templates" / "Live_Test_Short_Note_TEMPLATE.docx"
TEAM = "FTU — VeriTrade"
MAX_LISTED = 5          # rows named per bullet; a note is printed and signed, not scrolled

# Lines of the pipeline's own log that report something going wrong. Captured while the run
# streams (frontend/app.py) so section 4 quotes the run instead of summarising it from memory.
_ISSUE = re.compile(
    r"^\[(warn|error)\]"
    r"|^\[fetch\].*(fail|refus|timed? ?out|timeout|HTTP [45]\d\d|dropped|not allowed|skip)"
    r"|^\[extract\].*fail"
    r"|^\[portal\].*(refus|fail|timed? ?out)", re.I)


def run_issues(lines) -> list[str]:
    """The log lines that report a failure, de-duplicated, in the order they happened."""
    out: list[str] = []
    for ln in lines or []:
        s = str(ln).strip()
        if s and _ISSUE.search(s) and s not in out:
            out.append(s)
    return out


def machine_description() -> str:
    cpu = os.cpu_count() or 0
    return (f"{platform.node() or 'unknown host'} — {platform.system()} {platform.release()}, "
            f"{cpu} CPU threads, Python {platform.python_version()}")


# ── reading the runs ──────────────────────────────────────────────────────────────────
def _result(r):
    return (r or {}).get("result")


def _mappings(r) -> list:
    res = _result(r)
    return list(getattr(res, "mappings", []) or [])


def _meta(r):
    return getattr(_result(r), "meta", None)


def _real(ms) -> list:
    return [m for m in engine_compare.exported(ms)
            if (getattr(m, "law_name", "") or "") not in PLACEHOLDER_LAW_NAMES]


def _code(m) -> str:
    return to_rdtii_code(getattr(m, "indicator_id", "") or "")


def _cite(m) -> str:
    law = (getattr(m, "law_name", "") or "").strip()
    art = (getattr(m, "article_section", "") or "").strip()
    return f"{law}, {art}" if art and art != "(document)" else law


def _hosts(r) -> list[str]:
    meta = _meta(r)
    hs = []
    for d in getattr(meta, "documents", []) or []:
        h = urlparse(getattr(d, "source_url", "") or "").netloc.replace("www.", "")
        if h and h not in hs:
            hs.append(h)
    return hs


def _pillar_codes(pillar) -> list[str]:
    try:
        from ..rdtii.indicators import get_indicators          # noqa: PLC0415
        return [to_rdtii_code(i.indicator_id) for i in get_indicators(int(pillar))]
    except Exception:                                          # noqa: BLE001
        return []


def _n(k: int, one: str, many: str | None = None) -> str:
    return f"{k} {one if k == 1 else (many or one + 's')}"


def _more(total: int, shown: int) -> str:
    return f" (and {total - shown} more in the comparison file)" if total > shown else ""


# ── sections 3–5, drafted from the runs ───────────────────────────────────────────────
def auto_worked(state: dict) -> list[str]:
    a, b = state["runs"].get("A"), state["runs"].get("B")
    out: list[str] = []
    if a:
        hosts = _hosts(a)
        if a.get("documents"):
            out.append(
                f"Discovery from zero seed URLs: engine A's pass found "
                f"{_n(a['documents'], 'document')}"
                + (f" on {', '.join(hosts[:3])}" if hosts else "")
                + f" and fetched {a['fetched']} of them over the network; "
                f"{_n(a.get('provisions', 0), 'provision')} were extracted article by article.")
        ocr = [o for o in getattr(_meta(a), "ocr_reports", []) or []
               if getattr(o, "cer", None) is not None]
        if ocr and all(getattr(o, "cer_under_5pct", False) for o in ocr):
            worst = max(o.cer for o in ocr)
            out.append(f"OCR on scanned pages: {_n(len(ocr), 'document')} measured, worst "
                       f"character error rate {worst * 100:.2f}% (under the 5% bar).")
        rows = _real(_mappings(a))
        grounded = [m for m in rows
                    if getattr(getattr(m, "confidence", None), "snippet_grounding", 0) >= 0.99]
        if rows:
            out.append(f"Verbatim quotes: {len(grounded)} of {_n(len(rows), 'exported row')} "
                       f"from engine A quote words found in the statute text itself.")
    if b and b.get("fetched") == 0 and a and not b.get("failed") and not a.get("failed"):
        out.append(f"The second pass stayed off the network: engine B re-read engine A's "
                   f"{_n(b.get('documents', 0), 'document')} and fetched nothing, so the engine "
                   f"is the only thing that differs between the two passes.")
    s = (state.get("compare") or {}).get("summary")
    if s and s["both"]:
        out.append(f"Engine agreement: {s['identical']} of the {_n(s['both'], 'provision')} "
                   f"both engines exported carry the same indicator, citation and quote.")
    return out


def auto_broke(state: dict) -> list[str]:
    runs = state["runs"]
    out: list[str] = []
    for slot in ("A", "B"):
        r = runs.get(slot)
        if not r:
            continue
        meta = _meta(r)
        if r.get("failed"):
            out.append(f"Engine {slot}'s pass stopped with an error and exported nothing: "
                       f"{r['failed'][:240]}.")
        issues = list(r.get("issues") or [])
        if issues:
            shown = issues[:MAX_LISTED]
            out.append(f"Engine {slot}'s run log reported {_n(len(issues), 'problem')}: "
                       + "; ".join(i[:160] for i in shown) + _more(len(issues), len(shown))
                       + ".")
        if getattr(meta, "llm_provider", "") == "mock":
            out.append(f"Engine {slot} ran on the offline stand-in grader (mock), not a real "
                       f"model: its rows are keyword matches, not evidence.")
        incomplete = [m for m in _mappings(r)
                      if "Grading incomplete" in (getattr(m, "notes", "") or "")]
        if incomplete:
            out.append(f"Engine {slot}: grading did not finish for "
                       f"{_n(len({_code(m) for m in incomplete}), 'indicator')} "
                       f"({', '.join(sorted({_code(m) for m in incomplete}))}) — the model "
                       f"provider failed during the run.")
        if getattr(meta, "served_from_cache", None):
            out.append(f"Engine {slot}'s result was a stored one from "
                       f"{meta.served_from_cache}, not a run made in the hour.")
        if slot == "A" and r.get("documents", 0) and not r.get("fetched") \
                and not getattr(meta, "served_from_cache", None):
            out.append("Engine A fetched no document over the network during this pass: every "
                       "body came from the local cache, so this run did not exercise live "
                       "fetching.")
    b = runs.get("B")
    if b and b.get("fetched"):
        out.append(f"The second pass was not engine-isolated: engine B fetched "
                   f"{_n(b['fetched'], 'document')} of its own (the template requires 0), so the "
                   f"comparison mixes a change of engine with a change of documents.")
    a = runs.get("A")
    if (a and b and not b.get("failed") and a.get("documents")
            and b.get("documents", 0) < a.get("documents", 0)):
        out.append(f"Engine B read {b.get('documents', 0)} of engine A's {a['documents']} "
                   f"documents — the rest had no cached body.")
    if not out:
        out.append("No step failed: neither run's log reported a warning or an error, and every "
                   "document engine A found was fetched and read.")
    s = (state.get("compare") or {}).get("summary")
    if s and (s["indicator_differs"] or s["a_only"] or s["b_only"]):
        out.append(f"Indicator mapping where the engines part ways: "
                   f"{_n(s['indicator_differs'], 'provision')} mapped to different indicators, "
                   f"{s['a_only']} exported by engine A only and {s['b_only']} by engine B only. "
                   f"On those the choice of engine decides the answer.")
    return out


def auto_caution(state: dict) -> list[str]:
    runs = state["runs"]
    a, b = runs.get("A"), runs.get("B")
    out: list[str] = []
    pillar = (state.get("brief") or {}).get("pillar")
    codes = _pillar_codes(pillar) if pillar else []
    found = {_code(m) for r in (a, b) if r for m in _real(_mappings(r))}
    gaps = [c for c in codes if c not in found]
    if gaps and (a or b):
        out.append(f"Indicator{'s' if len(gaps) > 1 else ''} {', '.join(gaps)}: no provision "
                   f"exported by either engine. Check the portal by hand before recording "
                   f"'no such measure' — an empty cell can be a gap in what was reached.")
    cmp = (state.get("compare") or {}).get("rows") or []
    split = [r for r in cmp if r.found_by == engine_compare.BOTH
             and r.indicator == engine_compare.DIFFERS]
    if split:
        shown = split[:MAX_LISTED]
        out.append("Provisions the engines mapped to different indicators: " + "; ".join(
            f"{r.a.law}, {r.a.article or r.b.article} (A: {', '.join(r.a.indicators)} · "
            f"B: {', '.join(r.b.indicators)})" for r in shown) + _more(len(split), len(shown))
            + ".")
    lone = [r for r in cmp if r.found_by != engine_compare.BOTH]
    if lone:
        shown = lone[:MAX_LISTED]
        out.append("Found by one engine only: " + "; ".join(
            f"{(r.a or r.b).law}, {(r.a or r.b).article or '(document)'} → "
            f"{', '.join((r.a or r.b).indicators)} "
            f"({'A' if r.found_by == engine_compare.A_ONLY else 'B'})" for r in shown)
            + _more(len(lone), len(shown)) + ".")
    for slot, r in (("A", a), ("B", b)):
        if not r:
            continue
        weak = sorted((m for m in _real(_mappings(r))
                       if getattr(getattr(m, "review_status", ""), "value", "") == "pending_review"),
                      key=lambda m: getattr(m, "confidence_score", 0.0))
        if weak:
            shown = weak[:3]
            out.append(f"Engine {slot}'s least certain exported rows (below the 0.85 "
                       f"auto-accept line, not checked by a person): " + "; ".join(
                           f"{_code(m)} · {_cite(m)} "
                           f"({getattr(m, 'confidence_score', 0.0):.2f})" for m in shown)
                       + _more(len(weak), len(shown)) + ".")
    if not out and (a or b):
        out.append("Nothing in these runs singles out a row: both engines agree on every "
                   "exported provision and every row cleared the auto-accept line. Read the "
                   "quotes against the source links before relying on them.")
    return out


# ── the note's fields ─────────────────────────────────────────────────────────────────
def _engine_label(state: dict, slot: str) -> str:
    r = state["runs"].get(slot)
    if r:                                     # what the run actually used
        return f"{r.get('provider') or '—'} · {r.get('model') or '—'}"
    e = (state.get("engines") or {}).get(slot, {})
    return f"{e.get('provider', '—')} · {e.get('model', '—')} (not run)"


def _task(state: dict) -> str:
    b = state.get("brief") or {}
    said = (b.get("task") or "").strip()
    what = f"{b.get('economy', '—')} · pillar {b.get('pillar', '—')}"
    if not said:
        return what
    # The steward's words as read out; the picked economy/pillar only when the words omit them.
    return said if str(b.get("economy", "")).lower() in said.lower() else f"{said} ({what})"


def _bullets(items: list[str]) -> str:
    return "\n".join(f"• {i}" for i in items)


def fields(state: dict, *, now: datetime | None = None, machine: str | None = None) -> dict:
    """Everything the note prints, as plain values. Sections 3–5 take the operator's text when
    they wrote any, and the drafted account otherwise."""
    now = (now or datetime.now().astimezone())
    notes = state.get("notes") or {}
    runs = state["runs"]

    def fig(slot, key, fmt="{}"):
        r = runs.get(slot)
        if not r or r.get(key) is None:
            return "—"
        return fmt.format(r[key])

    by_hand = (notes.get("by_hand") or "").strip()
    return {
        "team": TEAM,
        "task": _task(state),
        "engine_a": _engine_label(state, "A"),
        "engine_b": _engine_label(state, "B"),
        "machine_time": f"{machine or machine_description()} · submitted "
                        f"{now.strftime('%Y-%m-%d %H:%M')} (UTC{now.strftime('%z')[:3]}:"
                        f"{now.strftime('%z')[3:]})",
        "exported": (fig("A", "exported"), fig("B", "exported")),
        "new": (fig("A", "new"), fig("B", "new")),
        "fetched": (fig("A", "fetched"), fig("B", "fetched")),
        "elapsed": (fig("A", "elapsed_min", "{:.1f}"), fig("B", "elapsed_min", "{:.1f}")),
        "cost": (fig("A", "cost_usd", "{:.4f}"), fig("B", "cost_usd", "{:.4f}")),
        "worked": (notes.get("worked") or "").strip() or _bullets(auto_worked(state)) or "—",
        "broke": (notes.get("broke") or "").strip() or _bullets(auto_broke(state)) or "—",
        "caution": (notes.get("caution") or "").strip() or _bullets(auto_caution(state)) or "—",
        "by_hand": by_hand,
    }


# ── outputs ───────────────────────────────────────────────────────────────────────────
def _fill(cell, text: str, size_pt: float = 9.5) -> None:
    """Write text into a template cell, keeping the cell's own first paragraph (and so its
    style); one paragraph per line."""
    import copy                                                 # noqa: PLC0415

    from docx.shared import Pt                                  # noqa: PLC0415
    from docx.text.paragraph import Paragraph                   # noqa: PLC0415
    # Sections 3 and 4 put their writing box in a table NESTED inside the cell; writing into
    # the outer cell lands the text under the empty box.
    while cell.tables:
        cell = cell.tables[0].rows[0].cells[0]
    lines = str(text).split("\n")
    target = next((p for p in cell.paragraphs if not p.text.strip()), None)
    if target is None:
        target = cell.add_paragraph()
    # The template's spare empty paragraphs would sit between our lines as blank gaps.
    for p in cell.paragraphs:
        if p is not target and not p.text.strip() and p._element is not target._element:
            p._element.getparent().remove(p._element)
    prev = target
    for i, line in enumerate(lines):
        if i == 0:
            p = target
        else:                                   # directly after the previous line
            new = copy.deepcopy(target._element)
            for r in new.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"):
                new.remove(r)
            prev._element.addnext(new)
            p = Paragraph(new, prev._parent)
        run = p.add_run(line)
        run.font.size = Pt(size_pt)
        prev = p


def to_docx(f: dict, template: Path = TEMPLATE) -> bytes:
    """The organisers' template with every field filled in. Signature cells stay empty."""
    from docx import Document                                   # noqa: PLC0415

    doc = Document(str(template))
    t_run, t_out, t_wb, t_caution, t_hand = doc.tables[:5]
    for row, key in zip(t_run.rows[1:], ("team", "task", "engine_a", "engine_b", "machine_time")):
        _fill(row.cells[1], f[key], 10)
    for row, key in zip(t_out.rows[1:], ("exported", "new", "fetched", "elapsed", "cost")):
        va, vb = f[key]
        _fill(row.cells[1], va, 10)
        cb = row.cells[2]
        if key == "fetched" and cb.text.strip():               # keep the printed "must be 0"
            cb.paragraphs[0].runs[0].text = f"{vb}   (must be 0)"
        else:
            _fill(cb, vb, 10)
    _fill(t_wb.rows[0].cells[0], f["worked"], 9)
    _fill(t_wb.rows[0].cells[2], f["broke"], 9)
    _fill(t_caution.rows[0].cells[0], f["caution"], 9)
    for p in doc.paragraphs:
        if p.text.lstrip().startswith("☐") and "typed in by hand" in p.text:
            # First box: nothing typed by hand. Second: something was. Tick exactly one.
            text = p.text
            boxes = [i for i, ch in enumerate(text) if ch == "☐"]
            i = boxes[1] if f["by_hand"] and len(boxes) > 1 else boxes[0]
            new = text[:i] + "☒" + text[i + 1:]
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
            p.runs[0].text = new
            break
    if f["by_hand"]:
        _fill(t_hand.rows[0].cells[0], f["by_hand"], 9.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_markdown(f: dict) -> str:
    """The same content as plain Markdown — the fallback when python-docx is missing."""
    rows = [("Provisions exported", "exported"), ("Of those, absent from the 2025 baseline", "new"),
            ("Documents fetched during this pass", "fetched"), ("Elapsed (minutes)", "elapsed"),
            ("Cost of this pass (US$)", "cost")]
    out = ["# Live test — short note", "Finale morning, 15 October 2026", "",
           "## 1 · The run", "", "| | |", "|---|---|",
           f"| Team name | {f['team']} |", f"| The task as read out | {f['task']} |",
           f"| Engine A, first pass — provider and model | {f['engine_a']} |",
           f"| Engine B, second pass — provider and model | {f['engine_b']} |",
           f"| Machine used, and time submitted | {f['machine_time']} |", "",
           "## 2 · What came out", "", "| | Engine A | Engine B |", "|---|---|---|"]
    out += [f"| {label} | {f[k][0]} | {f[k][1]} |" for label, k in rows]
    out += ["", "## 3 · What worked", "", f["worked"], "", "## 4 · What broke", "", f["broke"], "",
            "## 5 · What a reviewer should be cautious about", "", f["caution"], "",
            "## 6 · Anything done by hand", "",
            (f"☐ Nothing was typed in by hand. ☒ Something was — {f['by_hand']}" if f["by_hand"]
             else "☒ Nothing was typed in by hand. ☐ Something was"), "",
            "## 7 · Declaration", "",
            "Everything submitted is my team's own work, produced by the system frozen at our "
            "declared release tag, using only the engines declared on 30 September.", "",
            "Signed on behalf of the team: ________   Name and role: ________   "
            "Steward initials: ____", ""]
    return "\n".join(out)
